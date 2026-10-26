"""Per-upload VISTA pipeline for extracting, embedding, and indexing new reports."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
UPLOAD_ROOT = ROOT / "data" / "uploads"
RAW_ROOT = UPLOAD_ROOT / "raw"
JOB_ROOT = UPLOAD_ROOT / "jobs"
REGISTRY_PATH = UPLOAD_ROOT / "records.json"
DUCKDB_PATH = ROOT / "Database" / "DuckDB" / "data" / "metadata.duckdb"

PIPELINE_STEPS = [
    ("ingest", "Store original file", "Keep the source file unchanged in MinIO or local raw storage."),
    ("extract", "OCR, parse & clean", "Read report text or structured rows, then normalize the content."),
    ("fingerprint", "Build clinical fingerprint", "Pick out patient details, findings, and a search-ready summary."),
    ("embed", "Create 384-D vector", "Encode the clinical fingerprint for semantic comparison."),
    ("similarity", "Compare semantic meaning", "Score the new vector against the existing semantic box vectors."),
    ("box", "Assign semantic box", "Attach each new record to its nearest labeled box."),
    ("warehouse", "Write warehouse record", "Store the extracted record and vector metadata in DuckDB."),
    ("query", "Verify SQL access", "Run a parameterized count query against the new warehouse records."),
]
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".csv", ".tsv", ".json", ".jsonl", ".xlsx", ".txt", ".text", ".md", ".xml", ".html", ".htm", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
MAX_FILE_BYTES = 50 * 1024 * 1024
_jobs_lock = threading.RLock()
_pipeline_lock = threading.Lock()
_jobs: dict[str, dict[str, Any]] = {}
_records: list[dict[str, Any]] = []
_warehouse_manager = None
_dataset_import: dict[str, Any] = {
    "status": "idle", "processed": 0, "total": 0, "record_count": 0,
    "filename": "phase1_selection.txt", "error": None,
}


def _warehouse():
    global _warehouse_manager
    from config.settings import settings

    if not settings.NEON_DATABASE_URL:
        return None
    if _warehouse_manager is None:
        from Database.Neon.connection import NeonWarehouseManager
        _warehouse_manager = NeonWarehouseManager(settings.NEON_DATABASE_URL)
    return _warehouse_manager


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _new_steps() -> list[dict[str, Any]]:
    return [{"key": key, "label": label, "detail": detail, "status": "queued", "message": "Waiting"} for key, label, detail in PIPELINE_STEPS]


def _write_json_atomic(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _public_job(job: dict[str, Any]) -> dict[str, Any]:
    return json.loads(json.dumps(job, ensure_ascii=False))


def _save_job(job: dict[str, Any]) -> None:
    with _jobs_lock:
        job["updated_at"] = _now()
        _jobs[job["job_id"]] = job
        _write_json_atomic(JOB_ROOT / f"{job['job_id']}.json", job)
    warehouse = _warehouse()
    if warehouse:
        warehouse.upsert_job(job)


def _load_saved_state() -> None:
    global _records, _dataset_import
    try:
        payload = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
        _records = payload if isinstance(payload, list) else []
    except (OSError, json.JSONDecodeError):
        _records = []
    for path in sorted(JOB_ROOT.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)[:30]:
        try:
            job = json.loads(path.read_text(encoding="utf-8"))
            if job.get("job_id"):
                _jobs[job["job_id"]] = job
        except (OSError, json.JSONDecodeError):
            continue
    try:
        saved_import = json.loads((JOB_ROOT / "phase2_dataset_import.json").read_text(encoding="utf-8"))
        if isinstance(saved_import, dict):
            _dataset_import.update(saved_import)
    except (OSError, json.JSONDecodeError):
        pass


_load_saved_state()


def create_job(filename: str, content: bytes) -> dict[str, Any]:
    if not content:
        raise ValueError("The selected file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Files must be 50 MB or smaller.")
    safe_name = Path(filename.replace("\\", "/")).name.strip() or "medical-report"
    suffix = Path(safe_name).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Supported formats: PDF, DOCX, CSV/TSV, JSON/JSONL, XLSX, TXT/Markdown, XML/HTML, and common images.")
    job_id = uuid.uuid4().hex[:12]
    job = {
        "job_id": job_id,
        "filename": safe_name,
        "size_bytes": len(content),
        "status": "queued",
        "progress": 0,
        "current_step": "ingest",
        "created_at": _now(),
        "updated_at": _now(),
        "steps": _new_steps(),
        "result": None,
        "error": None,
    }
    _save_job(job)
    worker = threading.Thread(target=_run_pipeline, args=(job_id, safe_name, content), daemon=True)
    worker.start()
    return _public_job(job)


def get_job(job_id: str) -> dict[str, Any] | None:
    with _jobs_lock:
        job = _jobs.get(job_id)
        if job:
            return _public_job(job)
    warehouse = _warehouse()
    job = warehouse.get_job(job_id) if warehouse else None
    return _public_job(job) if job else None


def get_dataset_import() -> dict[str, Any]:
    with _jobs_lock:
        return _public_job(_dataset_import)


def start_dataset_import() -> dict[str, Any]:
    """Import the checked-in Phase 2 selection manifest from local disk."""
    global _dataset_import
    with _jobs_lock:
        if _dataset_import.get("status") in {"queued", "running"}:
            return _public_job(_dataset_import)
        if _dataset_import.get("status") == "complete":
            return _public_job(_dataset_import)
        manifest = ROOT / "phase 2_dataset" / "data" / "processed" / "phase1_selection.txt"
        if not manifest.is_file():
            raise ValueError("The Phase 2 selection manifest was not found in phase 2_dataset/data/processed.")
        entries = [line.strip() for line in manifest.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        selected_root = manifest.parent / "phase1_dataset"
        valid_entries = [entry for entry in entries if (selected_root / entry).is_file() and (selected_root / entry).suffix.lower() in ALLOWED_EXTENSIONS]
        if not valid_entries:
            raise ValueError("The Phase 2 selection manifest contains no readable report files.")
        _dataset_import = {
            "status": "queued", "processed": 0, "total": len(valid_entries), "record_count": 0,
            "filename": "phase1_selection.txt", "error": None,
        }
        worker = threading.Thread(target=_run_dataset_import, args=(selected_root, valid_entries), daemon=True, name="vista-phase2-import")
        worker.start()
        return _public_job(_dataset_import)


def _set_dataset_import(**updates: Any) -> None:
    with _jobs_lock:
        _dataset_import.update(updates)
        _write_json_atomic(JOB_ROOT / "phase2_dataset_import.json", _dataset_import)


def _run_dataset_import(selected_root: Path, entries: list[str]) -> None:
    """Extract the local selection manifest and embed it in moderate batches."""
    global _dataset_import
    from Backend.ETL.cleaner import MedicalTextCleaner
    from Backend.Embeddings.clinical_parser import ClinicalSectionParser
    from Backend.api import _engine

    _set_dataset_import(status="running", processed=0, record_count=0, started_at=_now())
    try:
        with _pipeline_lock:
            engine = _engine()
            existing_ids = {str(row.get("doc_id")) for row in uploaded_records()}
            prepared: list[dict[str, Any]] = []
            all_records: list[dict[str, Any]] = []
            vector_batch: list[str] = []
            id_hashes: set[str] = {
                MedicalTextCleaner.compute_hash(str(row.get("raw_text") or row.get("diagnostic_summary") or ""))
                for row in json.loads((ROOT / "data" / "processed" / "metadata.json").read_text(encoding="utf-8"))
                if row.get("raw_text") or row.get("diagnostic_summary")
            }
            batch_size = 128
            accepted = 0

            def flush() -> None:
                nonlocal accepted
                if not prepared:
                    return
                vectors = engine.embedder.encode(vector_batch, normalize_embeddings=True, show_progress_bar=False)
                if vectors.shape[1] != 384:
                    raise ValueError(f"Expected 384-D vectors, received {vectors.shape[1]} dimensions.")
                for record, vector in zip(prepared, vectors):
                    nearest = engine.faiss_index.search_box(vector, top_k=1)
                    match = nearest[0] if nearest else {"box_id": 0, "label": "General clinical records", "similarity_score": 0.0}
                    record["box_id"] = int(match["box_id"])
                    record["box_label"] = str(match.get("label", f"Clinical box {record['box_id'] + 1}"))
                    record["box_similarity"] = round(float(match.get("similarity_score", 0)), 4)
                    record["embedding"] = vector.astype(float).tolist()
                    record["priority_score"] = round(engine.priority_engine.compute_priority_score(1.0, 0.0, 0.65), 4)
                accepted += len(prepared)
                all_records.extend(prepared)
                _set_dataset_import(record_count=accepted)
                prepared.clear()
                vector_batch.clear()

            for index, relative in enumerate(entries, start=1):
                path = selected_root / relative
                digest = hashlib.sha1(relative.replace("\\", "/").casefold().encode("utf-8")).hexdigest()[:16]
                doc_id = f"phase2-{digest}"
                if doc_id in existing_ids:
                    _set_dataset_import(processed=index, record_count=accepted)
                    continue
                try:
                    if path.suffix.lower() in {".csv", ".tsv"}:
                        import pandas as pd
                        frame = pd.read_csv(path, dtype=str, sep="\t" if path.suffix.lower() == ".tsv" else ",", nrows=500).fillna("")
                        rows = frame.to_dict(orient="records")
                    else:
                        rows, _method = _read_source(path.name, path.read_bytes())
                    for row_index, source_row in enumerate(rows, start=1):
                        raw_text = source_row.get("raw_text") or source_row.get("text") or _flatten_row(source_row)
                        raw_text = MedicalTextCleaner.clean_text(str(raw_text))
                        if not raw_text:
                            continue
                        text_hash = MedicalTextCleaner.compute_hash(raw_text)
                        if text_hash in id_hashes:
                            continue
                        id_hashes.add(text_hash)
                        title = _pick(source_row, ("title", "report_title", "document_title", "name")) or path.stem
                        specialty = _pick(source_row, ("specialty", "category", "department", "diagnosis")) or "Clinical report"
                        summary = ClinicalSectionParser.build_diagnostic_summary(raw_text, title=title, category=specialty)
                        fingerprint = _fingerprint(raw_text)
                        row_patient = _pick(source_row, ("patient_name", "patient", "name", "full_name"))
                        if not row_patient:
                            first = _pick(source_row, ("firstname", "first_name"))
                            last = _pick(source_row, ("lastname", "last_name"))
                            row_patient = f"{first} {last}".strip()
                        if row_patient:
                            fingerprint["patient_name"] = row_patient
                        if fingerprint["patient_name"] == "Not found":
                            fingerprint["patient_name"] = _patient_from_filename(path.name) or "Not found"
                        if len(rows) > 1 and title == path.stem and fingerprint["patient_name"] != "Not found":
                            title = f"{fingerprint['patient_name']} · Medical report"
                        row_id = doc_id if len(rows) == 1 else f"{doc_id}-{row_index}"
                        prepared.append({
                            "doc_id": row_id, "title": title, "specialty": specialty, "raw_text": raw_text,
                            "full_text": raw_text, "snippet": raw_text[:300], "diagnostic_summary": summary,
                            "source_filename": f"Phase 2 dataset/{relative.replace(chr(92), '/')}", "job_id": "phase2-dataset",
                            "box_id": 0, "cosine_similarity": 0.0, "priority_score": 0.0, **fingerprint,
                        })
                        vector_batch.append(summary)
                        if len(prepared) >= batch_size:
                            flush()
                except Exception:
                    # A damaged or unsupported individual document should not stop the corpus import.
                    pass
                if index % batch_size == 0:
                    _set_dataset_import(processed=index, record_count=accepted)
            flush()

            records = all_records
            if records:
                _publish_records(records)
                import duckdb
                DUCKDB_PATH.parent.mkdir(parents=True, exist_ok=True)
                connection = duckdb.connect(str(DUCKDB_PATH))
                connection.execute("""CREATE TABLE IF NOT EXISTS uploaded_reports (
                    doc_id VARCHAR PRIMARY KEY, job_id VARCHAR, source_filename VARCHAR, title VARCHAR,
                    specialty VARCHAR, patient_name VARCHAR, age INTEGER, conditions VARCHAR, box_id INTEGER,
                    diagnostic_summary VARCHAR, raw_text VARCHAR, embedding_json VARCHAR, ingested_at TIMESTAMP
                )""")
                connection.executemany("""INSERT OR REPLACE INTO uploaded_reports
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, current_timestamp)""", [
                    (r["doc_id"], "phase2-dataset", r["source_filename"], r["title"], r["specialty"], r.get("patient_name", "Not found"), r.get("age"), json.dumps(r.get("conditions", [])), r["box_id"], r["diagnostic_summary"], r["raw_text"], json.dumps(r["embedding"])) for r in records
                ])
                connection.close()
            _set_dataset_import(status="complete", processed=len(entries), total=len(entries), record_count=len(records), completed_at=_now(), error=None)
    except Exception as exc:
        _set_dataset_import(status="failed", error=str(exc), failed_at=_now())


def list_jobs(limit: int = 20) -> list[dict[str, Any]]:
    with _jobs_lock:
        jobs = sorted(_jobs.values(), key=lambda item: item.get("created_at", ""), reverse=True)[:limit]
    warehouse = _warehouse()
    if warehouse:
        by_id = {job["job_id"]: job for job in warehouse.load_jobs(limit)}
        by_id.update({job["job_id"]: job for job in jobs})
        jobs = sorted(by_id.values(), key=lambda item: item.get("updated_at", ""), reverse=True)[:limit]
    return [_public_job(job) for job in jobs]


def uploaded_records(with_vectors: bool = False) -> list[dict[str, Any]]:
    with _jobs_lock:
        records = {record["doc_id"]: record for record in _records}
    warehouse = _warehouse()
    if warehouse:
        records.update({record["doc_id"]: record for record in warehouse.load_reports()})
    values = [_public_job(record) for record in records.values()]
    if with_vectors:
        return values
    return [{key: value for key, value in record.items() if key != "embedding"} for record in values]


def _publish_records(records: list[dict[str, Any]]) -> None:
    """Publish extracted reports immediately, then replace them as vectors finish."""
    with _jobs_lock:
        ids = {record["doc_id"] for record in records}
        retained = [record for record in _records if record["doc_id"] not in ids]
        retained.extend(records)
        warehouse = _warehouse()
        if warehouse:
            warehouse.upsert_reports(records)
        _write_json_atomic(REGISTRY_PATH, retained)
        _records[:] = retained


def _set_step(job: dict[str, Any], index: int, status: str, message: str) -> None:
    job["steps"][index]["status"] = status
    job["steps"][index]["message"] = message
    job["current_step"] = job["steps"][index]["key"]
    job["progress"] = round(index * 100 / len(PIPELINE_STEPS)) if status != "done" else round((index + 1) * 100 / len(PIPELINE_STEPS))
    job["status"] = "running"
    _save_job(job)


def _pick(mapping: dict[str, Any], names: tuple[str, ...]) -> str:
    lowered = {str(key).strip().casefold().replace(" ", "_"): value for key, value in mapping.items()}
    for name in names:
        value = lowered.get(name)
        if value is not None and str(value).strip() and str(value).strip().casefold() not in {"nan", "none", "null"}:
            return str(value).strip()
    return ""


def _flatten_row(row: Any) -> str:
    if isinstance(row, dict):
        return "\n".join(f"{key}: {value}" for key, value in row.items() if value is not None and str(value).strip())
    return str(row).strip()


def _read_source(filename: str, content: bytes) -> tuple[list[dict[str, Any]], str]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf).strip()
        if text:
            return [{"raw_text": text, "title": Path(filename).stem}], "Embedded PDF text"
        text = _run_paddle(content, suffix)
        return [{"raw_text": text, "title": Path(filename).stem}], "PaddleOCR · scanned PDF"
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return [{"raw_text": _run_paddle(content, suffix), "title": Path(filename).stem}], "PaddleOCR · image"
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(content))
        parts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return [{"raw_text": "\n".join(parts), "title": Path(filename).stem}], "DOCX text + tables"
    if suffix in {".csv", ".tsv", ".xlsx"}:
        import pandas as pd
        frame = pd.read_excel(io.BytesIO(content), dtype=str).fillna("") if suffix == ".xlsx" else pd.read_csv(io.BytesIO(content), dtype=str, sep="\t" if suffix == ".tsv" else ",", nrows=500).fillna("")
        rows = frame.to_dict(orient="records")
        if not rows:
            raise ValueError("No data rows were found in this file.")
        return rows, f"{suffix[1:].upper()} · first {len(frame):,} rows loaded"
    if suffix in {".json", ".jsonl"}:
        decoded = content.decode("utf-8-sig")
        if suffix == ".jsonl":
            items = [json.loads(line) for line in decoded.splitlines() if line.strip()]
        else:
            value = json.loads(decoded)
            items = value if isinstance(value, list) else next((value[key] for key in ("records", "data", "patients", "reports") if isinstance(value, dict) and isinstance(value.get(key), list)), [value])
        if not isinstance(items, list) or not items:
            raise ValueError("No report records were found in this JSON file.")
        return [{"record": item} if not isinstance(item, dict) else item for item in items[:500]], f"JSON · {len(items):,} records" + (" · first 500 queued" if len(items) > 500 else "")
    if suffix in {".txt", ".text", ".md"}:
        return [{"raw_text": content.decode("utf-8-sig", errors="replace"), "title": Path(filename).stem}], "Plain text"
    if suffix in {".xml", ".html", ".htm"}:
        from html.parser import HTMLParser
        class TextContent(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts: list[str] = []
            def handle_data(self, data: str):
                if data.strip():
                    self.parts.append(data.strip())
        parser = TextContent()
        parser.feed(content.decode("utf-8-sig", errors="replace"))
        return [{"raw_text": "\n".join(parser.parts), "title": Path(filename).stem}], f"{suffix[1:].upper()} text"
    raise ValueError("This file format is not supported yet.")


def _run_paddle(content: bytes, suffix: str) -> str:
    import sys
    ocr_root = str(ROOT / "project" / "code" / "vista_ocr")
    if ocr_root not in sys.path:
        sys.path.insert(0, ocr_root)
    from preprocessing.paddle_engine import extract_from_image_paddle, extract_from_scanned_pdf_paddle
    if suffix == ".pdf":
        return extract_from_scanned_pdf_paddle(content)
    return extract_from_image_paddle(content)


def _clean_patient_name(value: str) -> str:
    value = re.split(r"\s*(?:\||\b(?:lab\s*no\.?|age|gender|sex|reported|collected|reg(?:istration)?\s*date|report(?:ed)?\s*date|date\s*of\s*birth|dob|ref\.?\s*by|diabetes|thyroid|cardiac|risk|full\s*body|checkup|kidney|liver|function|anemia|workup|panel|clinical|medical|report)\b)", value, maxsplit=1, flags=re.IGNORECASE)[0]
    value = re.sub(r"^(?:mr|mrs|ms|miss|dr)\.?\s+", "", value.strip(), flags=re.IGNORECASE)
    value = re.sub(r"[^A-Za-z'’ -]", " ", value)
    words = value.split()
    return " ".join(words[:4]).strip()


def _patient_from_text(text: str) -> str:
    """Read names from both labeled clinical notes and common lab report layouts."""
    patterns = (
        r"\b(?:patient[\s_-]*name|name)\s*[:=]?\s*([A-Za-z][A-Za-z'’.-]*(?:\s+[A-Za-z][A-Za-z'’.-]*){1,3})",
        r"\b(?:mr|mrs|ms|miss|dr)\.?\s+([A-Za-z][A-Za-z'’.-]*(?:\s+[A-Za-z][A-Za-z'’.-]*){1,3})\s*(?::\s*){1,8}\s*(?:patient\s*)?name\b",
        r"\b(?:mr|mrs|ms|miss|dr)\.?\s+([A-Za-z][A-Za-z'’.-]*(?:\s+[A-Za-z][A-Za-z'’.-]*){1,3})(?=\s+(?:age|gender|sex|lab\s*no|reported|collected)\b)",
    )
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            name = _clean_patient_name(match.group(1))
            if len(name.split()) >= 2:
                return name
    return ""


def _fingerprint(text: str) -> dict[str, Any]:
    patient = re.search(r"(?im)\b(?:patient[\s_-]*(?:name)?|name)\s*[:=]\s*([^\n,;]{2,70})", text)
    age = re.search(r"(?i)\b(?:age\s*[:=]?\s*(\d{1,3})|(?:age\s*)?(\d{1,3})\s*(?:years?\s*old|y/?o))\b", text)
    named_age = re.search(r"(?i)\b([A-Z][A-Za-z'-]{1,39})\s+(?:age|aged)\s*(?:is\s*)?[:=]?\s*(\d{1,3})\b", text)
    if not patient:
        patient = re.search(r"(?im)\bpatient\s+([A-Z][A-Za-z'-]{1,39})(?=\s*[,;:]|\s+\d{1,3}\b)", text)
    if not age and named_age:
        age_value = named_age.group(2)
    else:
        age_value = next((part for part in age.groups() if part), "") if age else ""
    indicators = [
        ("diabetes", r"\b(?:diabet(?:es|ic)|high sugar|high blood sugar|hyperglyc(?:emia|aemia))\b"),
        ("hypertension", r"\b(?:hypertension|\bhtn\b|high bp|high blood pressure|elevated blood pressure)\b"),
        ("cardiac", r"\b(?:heart attack|myocardial infarction|troponin|cardiac)\b"),
        ("asthma", r"\basthma\b"),
        ("kidney", r"\b(?:kidney|renal)\b"),
        ("cancer", r"\b(?:cancer|carcinoma|tumou?r|oncology)\b"),
    ]
    conditions = [label for label, pattern in indicators if re.search(pattern, text, re.IGNORECASE)]
    bp = re.search(r"(?i)\b(?:bp|blood pressure)\s*[:=]?\s*(\d{2,3})\s*/\s*(\d{2,3})\b", text)
    glucose = re.search(r"(?i)\b(?:glucose|blood sugar|hba1c)\s*[:=]?\s*(\d+(?:\.\d+)?)\s*(mg\s*/\s*dL|mmol\s*/\s*L|%)?", text)
    vitals = {}
    if bp:
        systolic, diastolic = int(bp.group(1)), int(bp.group(2))
        vitals["BP"] = f"{systolic}/{diastolic}"
        if (systolic >= 140 or diastolic >= 90) and "hypertension" not in conditions:
            conditions.append("hypertension")
    if glucose:
        vitals["Glucose"] = f"{glucose.group(1)} {glucose.group(2) or ''}".strip()
    patient_name = _patient_from_text(text)
    if not patient_name and patient:
        patient_name = _clean_patient_name(patient.group(1))
    if not patient_name and named_age:
        patient_name = _clean_patient_name(named_age.group(1))
    return {"patient_name": patient_name or "Not found", "age": int(age_value) if age_value else None, "conditions": conditions or ["No common finding detected"], "vitals": vitals}


def _patient_from_filename(filename: str) -> str:
    """Use a descriptive report filename as a patient-name fallback."""
    stem = Path(filename).stem
    stem = re.sub(r"(?:[_-](?:M|F|MALE|FEMALE))?[_-]\d{1,3}[_-]Yrs?\b.*$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"[_-](?:\d{8,}|\d{4}[-_]\d{2}[-_]\d{2}).*$", "", stem)
    stem = re.sub(r"\b(?:report|medical|patient|scan|final|copy|diabetes|thyroid|cardiac|risk|full|body|checkup|kidney|liver|function|anemia|workup|panel|clinical)\b.*$", " ", stem, flags=re.IGNORECASE)
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", stem.replace("_", " ").replace("-", " "))
    if words and words[0].casefold() in {"mr", "mrs", "ms", "miss", "dr"}:
        words = words[1:]
    return " ".join(words).strip() if len(words) >= 2 else ""


def _run_pipeline(job_id: str, filename: str, content: bytes) -> None:
    from Backend.ETL.cleaner import MedicalTextCleaner
    from Backend.Embeddings.clinical_parser import ClinicalSectionParser
    from config.settings import settings

    job = _jobs[job_id]
    records: list[dict[str, Any]] = []
    try:
        with _pipeline_lock:
            _set_step(job, 0, "running", "Saving the unchanged source file")
            safe_name = Path(filename).name
            raw_path = RAW_ROOT / job_id / safe_name
            raw_path.parent.mkdir(parents=True, exist_ok=True)
            raw_path.write_bytes(content)
            raw_object = None
            endpoint = settings.MINIO_ENDPOINT.replace("http://", "").replace("https://", "")
            if endpoint not in {"localhost:9000", "127.0.0.1:9000"}:
                try:
                    from Database.MinIO.client import MinIOStorageManager
                    storage = MinIOStorageManager(endpoint=endpoint, access_key=settings.MINIO_ACCESS_KEY, secret_key=settings.MINIO_SECRET_KEY, secure=settings.MINIO_SECURE, bucket_name=settings.MINIO_BUCKET_NAME)
                    raw_object = storage.upload_file(raw_path, f"uploads/raw/{job_id}/{safe_name}")
                    raw_storage = f"MinIO ? {settings.MINIO_BUCKET_NAME}/{raw_object}"
                except Exception:
                    raw_storage = f"Local raw layer ? data/uploads/raw/{job_id}/{safe_name}"
            else:
                raw_storage = f"Local raw layer ? data/uploads/raw/{job_id}/{safe_name}"
            _set_step(job, 0, "done", raw_storage)

            _set_step(job, 1, "running", "Detecting file type and extracting content")
            source_rows, extraction_method = _read_source(filename, content)
            prepared = []
            seen_hashes = {MedicalTextCleaner.compute_hash(str(row.get("raw_text", ""))) for row in uploaded_records()}
            for index, source_row in enumerate(source_rows, start=1):
                raw_text = source_row.get("raw_text") or source_row.get("text") or _flatten_row(source_row)
                raw_text = MedicalTextCleaner.clean_text(str(raw_text))
                if not raw_text:
                    continue
                content_hash = MedicalTextCleaner.compute_hash(raw_text)
                if content_hash in seen_hashes:
                    continue
                seen_hashes.add(content_hash)
                title = _pick(source_row, ("title", "report_title", "document_title", "name")) or f"{Path(filename).stem} · record {index}"
                specialty = _pick(source_row, ("specialty", "category", "department", "diagnosis")) or "Clinical report"
                summary = ClinicalSectionParser.build_diagnostic_summary(raw_text, title=title, category=specialty)
                fingerprint = _fingerprint(raw_text)
                if fingerprint["patient_name"] == "Not found":
                    fingerprint["patient_name"] = _patient_from_filename(filename) or "Not found"
                prepared.append({"doc_id": f"upload-{job_id}-{index:03d}", "title": title, "specialty": specialty, "raw_text": raw_text, "full_text": raw_text, "snippet": raw_text[:300], "diagnostic_summary": summary, "source_filename": filename, "job_id": job_id, "box_id": 0, "cosine_similarity": 0.0, "priority_score": 0.0, **fingerprint})
            if not prepared:
                raise ValueError("No readable clinical text was extracted from this file.")
            records = prepared
            # Make extracted reports searchable while vectorization and
            # warehouse persistence continue in the background.
            _publish_records(records)
            _set_step(job, 1, "done", f"{extraction_method} · {len(records):,} unique record(s) · cleaned and deduplicated")

            _set_step(job, 2, "running", "Extracting patient details and clinical indicators")
            _set_step(job, 2, "done", f"Built {len(records):,} fingerprint(s) from clinical summary, patient, age, and findings")

            _set_step(job, 3, "running", "Loading the VISTA sentence-transformer")
            from Backend.api import _engine
            engine = _engine()
            vectors = engine.embedder.encode([record["diagnostic_summary"] for record in records], normalize_embeddings=True, show_progress_bar=False)
            if vectors.shape[1] != 384:
                raise ValueError(f"Expected 384-D vectors, received {vectors.shape[1]} dimensions.")
            _set_step(job, 3, "done", f"{len(records):,} normalized vectors · {vectors.shape[1]} dimensions")

            _set_step(job, 4, "running", "Calculating cosine similarity against box fingerprints")
            for record, vector in zip(records, vectors):
                matches = engine.faiss_index.search_box(vector, top_k=1)
                if not matches:
                    raise ValueError("The VISTA semantic box index has no available clusters.")
                match = matches[0]
                record["matched_box_id"] = int(match["box_id"])
                record["matched_box_label"] = str(match.get("label", f"Clinical box {record['matched_box_id'] + 1}"))
                record["box_similarity"] = round(float(match.get("similarity_score", 0)), 4)
            _set_step(job, 4, "done", "Cosine similarity computed against the indexed box fingerprints")

            _set_step(job, 5, "running", "Assigning each fingerprint to its nearest labeled box")
            for record, vector in zip(records, vectors):
                record["box_id"] = record.pop("matched_box_id")
                record["box_label"] = record.pop("matched_box_label")
                record["embedding"] = vector.astype(float).tolist()
                record["priority_score"] = round(engine.priority_engine.compute_priority_score(1.0, 0.0, 0.65), 4)
            _set_step(job, 5, "done", "Raw record and vector fingerprint linked to the labeled semantic box")

            _set_step(job, 6, "running", "Writing searchable records to DuckDB")
            import duckdb
            DUCKDB_PATH.parent.mkdir(parents=True, exist_ok=True)
            connection = duckdb.connect(str(DUCKDB_PATH))
            connection.execute("""CREATE TABLE IF NOT EXISTS uploaded_reports (
                doc_id VARCHAR PRIMARY KEY, job_id VARCHAR, source_filename VARCHAR, title VARCHAR,
                specialty VARCHAR, patient_name VARCHAR, age INTEGER, conditions VARCHAR, box_id INTEGER,
                diagnostic_summary VARCHAR, raw_text VARCHAR, embedding_json VARCHAR, ingested_at TIMESTAMP
            )""")
            connection.executemany("""INSERT OR REPLACE INTO uploaded_reports
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, current_timestamp)""", [
                (r["doc_id"], job_id, filename, r["title"], r["specialty"], r["patient_name"], r["age"], json.dumps(r["conditions"]), r["box_id"], r["diagnostic_summary"], r["raw_text"], json.dumps(r["embedding"])) for r in records
            ])
            connection.close()
            _publish_records(records)
            _set_step(job, 6, "done", f"{len(records):,} record(s) indexed in local DuckDB metadata warehouse")

            _set_step(job, 7, "running", "Checking that the warehouse can answer record counts")
            connection = duckdb.connect(str(DUCKDB_PATH), read_only=True)
            count = connection.execute("SELECT COUNT(*) FROM uploaded_reports WHERE job_id = ?", [job_id]).fetchone()[0]
            connection.close()
            _set_step(job, 7, "done", f"SQL count returned {count:,} uploaded record(s)")

            profiles = [{key: record[key] for key in ("doc_id", "title", "specialty", "patient_name", "age", "conditions", "vitals", "box_id", "box_label", "box_similarity", "snippet")} for record in records[:20]]
            job["status"] = "complete"
            job["progress"] = 100
            job["current_step"] = None
            job["result"] = {"record_count": len(records), "extraction_method": extraction_method, "embedding_dimension": int(vectors.shape[1]), "storage": raw_storage, "warehouse": "DuckDB", "sql_count": int(count), "profiles": profiles}
            job["completed_at"] = _now()
            _save_job(job)
    except Exception as exc:
        job["status"] = "failed"
        job["error"] = str(exc)
        for step in job["steps"]:
            if step["status"] == "running":
                step["status"] = "failed"
                step["message"] = str(exc)
            elif step["status"] == "queued":
                step["status"] = "skipped"
                step["message"] = "Stopped after an earlier step failed"
        _save_job(job)


def execute_count_query(question: str) -> dict[str, Any]:
    """Translate a count-style question to a parameterized DuckDB query."""
    import duckdb
    from Backend.api import _documents, _box_data

    docs = _documents()
    connection = duckdb.connect(":memory:")
    connection.execute("CREATE TABLE vista_reports (doc_id VARCHAR, title VARCHAR, specialty VARCHAR, box_id INTEGER, full_text VARCHAR)")
    connection.executemany("INSERT INTO vista_reports VALUES (?, ?, ?, ?, ?)", [(d["doc_id"], d["title"], d["specialty"], d["box_id"], d.get("full_text", "")) for d in docs])
    normalized = question.casefold()
    box = next((item for item in _box_data() if item["label"].casefold() in normalized), None)
    search_terms = [term for term in re.findall(r"[a-z][a-z-]{3,}", normalized) if term not in {"many", "reports", "records", "patients", "documents", "mention", "mentions", "contain", "contains", "with", "show", "count", "total", "from", "have", "does", "there", "what", "number", "that", "this", "the", "how"}]
    if box:
        sql = "SELECT COUNT(*) AS report_count FROM vista_reports WHERE box_id = ?"
        params: list[Any] = [int(box["box_id"])]
        plan = f"Count reports assigned to {box['label']}"
    elif search_terms:
        term = max(search_terms, key=len)
        sql = "SELECT COUNT(*) AS report_count FROM vista_reports WHERE lower(concat_ws(' ', title, specialty, full_text)) LIKE ?"
        params = [f"%{term}%"]
        plan = f"Count reports whose clinical text matches “{term}”"
    else:
        sql = "SELECT COUNT(*) AS report_count FROM vista_reports"
        params = []
        plan = "Count all reports in the searchable corpus"
    result = int(connection.execute(sql, params).fetchone()[0])
    connection.close()
    return {"question": question, "sql": sql, "result": result, "plan": plan, "source": "DuckDB · searchable corpus"}
